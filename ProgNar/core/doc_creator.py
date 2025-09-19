import json
from docx import Document
from docx.shared import Inches
from tkinter import filedialog
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm,Pt, RGBColor


def create_document():
    with open("data/temp_cart.json", "r", encoding="utf-8") as f:
        data = json.load(f)
    items = data["items"]  # to jest lista narzędzi

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)  # szerokość A4
    section.page_height = Mm(210)  # wysokość A4
    section.left_margin = Mm(5)
    section.right_margin = Mm(5)
    doc.add_heading("Wycena narzędziowa", level=1)
    doc.add_paragraph("Data: [tu możesz dodać datę]")

    table = doc.add_table(rows=1, cols=14)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'L.P'
    hdr[1].text = 'Nazwa'
    hdr[2].text = 'Ø OD'
    hdr[3].text = 'Ø Trzonka'
    hdr[4].text = 'L [mm]'
    hdr[5].text = 'Ilość ostrzy'
    hdr[6].text = 'Ilość sztuk'
    hdr[7].text = 'Cena regeneracji / szt'
    hdr[8].text = 'Wartość regeneracji'
    hdr[9].text = 'Powłoka'
    hdr[10].text = 'Nazwa powłoki'
    hdr[11].text = 'Cena powlekania'
    hdr[12].text = 'Wartość powlekania'
    hdr[13].text = 'Uwagi'

    #FONT
    for cell in hdr:
        for paragraph in cell.paragraphs:
            run = paragraph.runs[0]
            run.font.name = 'Calibri'  # czcionka
            run.font.size = Pt(10)  # rozmiar
            run.font.bold = True  # pogrubienie
            run.font.color.rgb = RGBColor(0, 0, 0)  # kolor czarny
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # for item in items:
    #     row = table.add_row().cells
    #     row[0].text = item.get("Nazwa", "")
    #     row[1].text = item.get("Srednica", "")
    #     row[2].text = item.get("Ilosc sztuk", "")
    #     row[3].text = item.get("Powloka", "")
    #     row[4].text = item.get("Cena szlifowania", "")
    #     row[5].text = item.get("Cena powlekania", "")

    save_path = filedialog.asksaveasfilename(
        defaultextension=".docx",
        filetypes=[("Dokument Word", "*.docx")],
        title="Zapisz raport jako"
    )
    if not save_path:
        print("Zapis anulowany.")
        return
    doc.save(save_path)
    print(f"Raport zapisany: {save_path}")